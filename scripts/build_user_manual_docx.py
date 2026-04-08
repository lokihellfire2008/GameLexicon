from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "GameLexicon_User_Manual_v1.0.md"
OUTPUT = ROOT / "docs" / "GameLexicon_User_Manual_v1.0.docx"


@dataclass
class ScreenshotBlock:
    title: str
    recommendation: str
    purpose: str


def add_toc(paragraph):
    run = paragraph.add_run()

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'

    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")

    text = OxmlElement("w:t")
    text.text = "Right-click and update the field in Word to refresh the table of contents."
    fld_separate.append(text)

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_separate)
    run._r.append(fld_end)


def ensure_styles(document: Document):
    styles = document.styles

    if "Manual Bullet" not in styles:
        style = styles.add_style("Manual Bullet", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = styles["List Bullet"]
        style.font.name = "Aptos"
        style.font.size = Pt(10.5)

    if "Manual Number" not in styles:
        style = styles.add_style("Manual Number", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = styles["List Number"]
        style.font.name = "Aptos"
        style.font.size = Pt(10.5)


def apply_document_defaults(document: Document):
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size in [("Title", 22), ("Heading 1", 16), ("Heading 2", 13)]:
        style = document.styles[style_name]
        style.font.name = "Aptos"
        style.font.size = Pt(size)


def parse_screenshot(line: str) -> ScreenshotBlock | None:
    match = re.fullmatch(r"\[\[SCREENSHOT:\s*(.*?)\|(.*?)\|(.*?)\]\]", line.strip())
    if not match:
        return None
    return ScreenshotBlock(
        title=match.group(1).strip(),
        recommendation=match.group(2).strip(),
        purpose=match.group(3).strip(),
    )


def is_table_line(line: str) -> bool:
    stripped = line.strip()
    return stripped.startswith("|") and stripped.endswith("|")


def split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def set_cell_text(cell, text: str, bold: bool = False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(10)
    run.bold = bold


def add_screenshot_box(document: Document, block: ScreenshotBlock):
    table = document.add_table(rows=3, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    rows = [
        ("Screenshot Placeholder", block.title),
        ("Recommended Capture", block.recommendation),
        ("Purpose", block.purpose),
    ]

    for idx, (label, value) in enumerate(rows):
        cell = table.rows[idx].cells[0]
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        label_run = paragraph.add_run(f"{label}: ")
        label_run.bold = True
        label_run.font.name = "Aptos"
        label_run.font.size = Pt(10.5)
        value_run = paragraph.add_run(value)
        value_run.font.name = "Aptos"
        value_run.font.size = Pt(10.5)

    document.add_paragraph("")


def add_table(document: Document, lines: list[str]):
    rows = [split_table_row(line) for line in lines]
    if len(rows) < 2:
        return

    header = rows[0]
    body = []
    for row in rows[1:]:
        compact = [cell.replace(" ", "") for cell in row]
        if compact and all(re.fullmatch(r":?-{3,}:?", cell) for cell in compact):
            continue
        body.append(row)

    table = document.add_table(rows=1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    for idx, text in enumerate(header):
        set_cell_text(table.rows[0].cells[idx], text, bold=True)

    for row in body:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            set_cell_text(cells[idx], text)

    document.add_paragraph("")


def add_footer_page_number(document: Document):
    for section in document.sections:
        paragraph = section.footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        title_run = paragraph.add_run("GameLexicon User Manual  ")
        title_run.font.name = "Aptos"
        title_run.font.size = Pt(9)

        label_run = paragraph.add_run("Page ")
        label_run.font.name = "Aptos"
        label_run.font.size = Pt(9)

        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = "PAGE"
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")

        page_run = paragraph.add_run()
        page_run._r.append(fld_begin)
        page_run._r.append(instr)
        page_run._r.append(fld_end)


def enable_update_fields_on_open(document: Document):
    settings = document.settings.element
    node = OxmlElement("w:updateFields")
    node.set(qn("w:val"), "true")
    settings.append(node)


def build_document():
    if not SOURCE.exists():
        raise FileNotFoundError(f"Manual source not found: {SOURCE}")

    document = Document()
    ensure_styles(document)
    apply_document_defaults(document)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    paragraph_buffer: list[str] = []

    def flush_paragraphs():
        nonlocal paragraph_buffer
        if not paragraph_buffer:
            return
        text = " ".join(part.strip() for part in paragraph_buffer).strip()
        if text:
            paragraph = document.add_paragraph(text)
            paragraph.paragraph_format.space_after = Pt(6)
        paragraph_buffer = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped == "":
            flush_paragraphs()
            i += 1
            continue

        if stripped == "[TOC]":
            flush_paragraphs()
            toc_heading = document.add_heading("Table of Contents", level=1)
            toc_heading.paragraph_format.space_after = Pt(6)
            add_toc(document.add_paragraph())
            document.add_paragraph("")
            i += 1
            continue

        screenshot = parse_screenshot(stripped)
        if screenshot:
            flush_paragraphs()
            add_screenshot_box(document, screenshot)
            i += 1
            continue

        if is_table_line(stripped):
            flush_paragraphs()
            table_lines = []
            while i < len(lines) and is_table_line(lines[i].strip()):
                table_lines.append(lines[i].strip())
                i += 1
            add_table(document, table_lines)
            continue

        if stripped.startswith("# "):
            flush_paragraphs()
            heading = document.add_heading(stripped[2:].strip(), level=0)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        if stripped.startswith("## "):
            flush_paragraphs()
            document.add_heading(stripped[3:].strip(), level=1)
            i += 1
            continue

        if stripped.startswith("### "):
            flush_paragraphs()
            document.add_heading(stripped[4:].strip(), level=2)
            i += 1
            continue

        if stripped.startswith("- "):
            flush_paragraphs()
            paragraph = document.add_paragraph(stripped[2:].strip(), style="Manual Bullet")
            paragraph.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            flush_paragraphs()
            text = re.sub(r"^\d+\.\s+", "", stripped)
            paragraph = document.add_paragraph(text, style="Manual Number")
            paragraph.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        paragraph_buffer.append(stripped)
        i += 1

    flush_paragraphs()
    enable_update_fields_on_open(document)
    add_footer_page_number(document)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)


if __name__ == "__main__":
    build_document()
    print(OUTPUT)
